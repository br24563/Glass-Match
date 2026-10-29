"""Manufacturers' Excel data workbooks -> extra GlassMatch properties.

The Zemax .agf files carry n_d, V_d, dispersion, CTE, density, transmission
and dn/dT, but none of the thermal / mechanical figures a mount or athermal
design needs. SCHOTT, OHARA and CDGM publish those in their own Excel data
workbooks; this module adds them to the glasses the maker's .agf already
defines. It never creates a glass, never touches the .agf's n_d / V_d (it
only reports where the workbook disagrees), and never stores a property the
glass already has from its .agf or an earlier workbook (it reports a
disagreeing value instead).

Like glassmatch/importers/hikari_xlsx.py, columns are found by their header
text, never by position, so a re-ordered workbook fails loudly instead of
mis-assigning values. Units come from the workbooks themselves:

* SCHOTT - its "Description" sheet states each unit ("Young's modolus [GPa]",
  "heat conductivity [W/mK]", "stress optical coefficient in [10-6MPa-1]",
  ...); the importer refuses the workbook if those texts change. No shear
  modulus is given, so E = 2G(1 + nu) cannot be checked.
* CDGM - the header carries the unit: E(10^7 Pa), G(10^7 Pa), HK(10^7 Pa),
  B(10^-12/Pa), lambda[W/(m K)], cp[J/(g K)]. E and G are converted to GPa
  (x 0.01) and E = 2G(1 + mu) is checked on every row. Knoop hardness stays
  in CDGM's printed unit (1e7 Pa), not converted to an HK number.
* HOYA - the codes row (Tg, lambda, Cp, E, G, mu, B) carries no unit, but
  the workbook's own "glass_data_sheet" prints them (Tg(C), lambda[W/(m K)],
  Cp[kJ/(kg K)] = J/(g K), E(GPa), G(GPa), B(10^-12/Pa)); the importer
  refuses the workbook if those texts are missing. E = 2G(1 + mu) is checked
  on every row. The current HOYA file could not be downloaded (2026-10-28);
  this reads HOYA's 2026-04-01 workbook as redistributed by the open-source
  opticalglass package, recorded under a -MIRROR source_id.
* OHARA - the header names E, G and nu with no unit. E = 2G(1 + sigma) on
  every row pins E and G to one unit, and its magnitude (median E of an oxide
  glass line, 40-150) to GPa rather than 1e8 N/m^2 or kgf/mm^2. beta is in
  nm/cm/1e5 Pa (OHARA's catalog notes, "Mechanical properties"), which is
  numerically 1e-12/Pa.

Codes whose meaning the workbook does not state are left out rather than
guessed: CDGM's RC(S) and RA(S) surface-method classes (CDGM's CR column is
the .agf climate class verbatim, so it adds nothing), and HOYA's DA, Tblue,
DNaOH, DSTPP, D0 and DH, whose titles in the row above the codes do not line
up with them (DA sits under "Staining Resistance (powder method)").
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from statistics import median

import pandas as pd

# |workbook - .agf| beyond these is reported. Workbooks print n_d to 5
# decimals and V_d to 2; the .agf carries more digits.
ND_TOL = 2e-5
VD_TOL = 0.02
# E = 2G(1 + nu) with E and G printed to ~3 figures and nu to 3 decimals.
MODULI_TOL = 0.03
# More rows than this failing the moduli check means a wrong unit or column.
MODULI_MAX_FAIL = 0.2


def _spec(header, prop, unit="", scale=1.0, note="", sheet_group=None):
    """One column: exact header text (space- and dot-insensitive)."""
    return {"header": header, "prop": prop, "unit": unit, "scale": scale,
            "note": note or header, "group": sheet_group}


# Per layout: which sheets, how to find the header row, and the columns.
LAYOUTS = {
    "SCHOTT": {
        "maker": "SCHOTT", "sheets": ("Preferred glasses", "Inquiry glasses"),
        "name_header": "Glass",
        "nd": "nd", "vd": "vd",
        "columns": [
            _spec("Tg", "tg", "C", note="T_g (ISO 7884-8)"),
            _spec("Heat capacity (cp)", "specific_heat", "J/(g K)",
                  note="heat capacity cp"),
            _spec("Heat conductivity(lambda)", "thermal_conductivity",
                  "W/(m K)", note="heat conductivity"),
            _spec("Young's modulus (E)", "youngs_modulus", "GPa",
                  note="Young's modulus"),
            _spec("Poisson ratio (µ)", "poisson_ratio", note="Poisson ratio"),
            _spec("Knoop hardness (HK)", "knoop_hardness", "HK",
                  note="Knoop hardness (ISO 9385)"),
            _spec("K", "stress_optic_coefficient",
                  "1e-12/Pa", note="stress optical coefficient K "
                                   "(catalog 1e-6/MPa)"),
            _spec("CR", "climate_resistance", "class",
                  note="climatic resistance class CR"),
            _spec("FR", "stain_resistance", "class",
                  note="stain resistance class FR"),
            _spec("SR", "acid_resistance", "class",
                  note="acid resistance class SR (ISO 8424)"),
            _spec("AR", "alkali_resistance", "class",
                  note="alkali resistance class AR (ISO 10629)"),
            _spec("PR", "phosphate_resistance", "class",
                  note="phosphate resistance class PR (ISO 9689)"),
            _spec("WR-J", "water_resistance_powder", "class",
                  note="water resistance class, JOGIS (powder method)"),
            _spec("SR-J", "acid_resistance_powder", "class",
                  note="acid resistance class, JOGIS (powder method)"),
        ],
        # Description-sheet code -> text that must appear in its description.
        "unit_texts": {
            "Tg": "[°C]", "Heat capacity (cp)": "[J/gK]",
            "Heat conductivity(lambda)": "[W/mK]",
            "Young's modulus (E)": "[GPa]",
            "K": "[10-6MPa-1]",
        },
    },
    "CDGM": {
        "maker": "CDGM",
        "sheets": ("optical glass", "i line glass", "infrared glass"),
        "name_header": "Glass", "end_marker": "over!",
        "nd": "nd", "vd": "νd",
        "columns": [
            _spec("Tg(℃)", "tg", "C", note="T_g"),
            _spec("λ[W/(m·K)]", "thermal_conductivity", "W/(m K)",
                  note="thermal conductivity"),
            _spec("cp[J/(g·K)]", "specific_heat", "J/(g K)",
                  note="specific heat"),
            _spec("HK(107Pa)", "knoop_hardness", "1e7 Pa",
                  note="Knoop hardness HK, as printed in 1e7 Pa"),
            _spec("E(107Pa)", "youngs_modulus", "GPa", 0.01,
                  note="Young's modulus (catalog 1e7 Pa)"),
            _spec("G(107Pa)", "shear_modulus", "GPa", 0.01,
                  note="shear modulus (catalog 1e7 Pa)"),
            _spec("μ", "poisson_ratio", note="Poisson's ratio"),
            _spec("B(10-12/Pa)", "stress_optic_coefficient", "1e-12/Pa",
                  note="stress-optic coefficient B"),
            _spec("DW", "water_resistance_powder", "class",
                  note="water resistance DW (powder method)"),
            _spec("DA", "acid_resistance_powder", "class",
                  note="acid resistance DA (powder method)"),
        ],
        "moduli": True,
    },
    "HOYA": {
        "maker": "HOYA", "sheet_pattern": r"\d{8}data_table1",
        "name_header": "Glass Type", "header_depth": 4,
        "nd": "nd", "vd": "νd",
        "columns": [
            _spec("Tg", "tg", "C", note="transition temperature Tg"),
            _spec("λ", "thermal_conductivity", "W/(m K)",
                  note="thermal conductivity"),
            _spec("Cp", "specific_heat", "J/(g K)",
                  note="specific heat capacity (catalog kJ/(kg K))"),
            _spec("HK", "knoop_hardness", "HK", note="Knoop hardness"),
            _spec("E", "youngs_modulus", "GPa", note="Young's modulus"),
            _spec("G", "shear_modulus", "GPa", note="rigidity modulus"),
            _spec("μ", "poisson_ratio", note="Poisson's ratio"),
            _spec("B", "stress_optic_coefficient", "1e-12/Pa",
                  note="photoelastic constant B"),
            _spec("DW", "water_resistance_powder", "class",
                  note="water durability DW (powder method)"),
        ],
        "unit_sheet": "glass_data_sheet",
        "unit_strings": ("Tg(℃)", "λ[W/(m･K)]", "CP[kJ/(kg･K)]", "E(GPa)",
                         "G(GPa)", "B(10-12/Pa)"),
        "moduli": True,
    },
    "OHARA": {
        "maker": "OHARA", "sheets": None, "name_header": "Glass",
        "nd": "nd", "vd": "νd",
        "columns": [
            _spec("Tg（℃）", "tg", "C", note="T_g"),
            _spec("λ(W/(m･K))", "thermal_conductivity", "W/(m K)",
                  note="thermal conductivity"),
            _spec("Young's (E)", "youngs_modulus", "GPa",
                  note="Young's modulus E"),
            _spec("Rigidity (G)", "shear_modulus", "GPa",
                  note="modulus of rigidity G"),
            _spec("Poisson's(σ)", "poisson_ratio", note="Poisson's ratio"),
            _spec("Knoop (Hk)", "knoop_hardness", "HK", note="Knoop hardness"),
            _spec("β", "stress_optic_coefficient", "1e-12/Pa",
                  note="photoelastic constant beta (catalog nm/cm/1e5 Pa)"),
            _spec("RW(P)", "water_resistance_powder", "class",
                  note="water resistance RW(P) (powder method)"),
            _spec("RA(P)", "acid_resistance_powder", "class",
                  note="acid resistance RA(P) (powder method)"),
            _spec("W(S)max", "weathering_resistance_surface_max", "class",
                  note="weathering resistance W(S), surface method, max"),
            _spec("W(S)min", "weathering_resistance_surface_min", "class",
                  note="weathering resistance W(S), surface method, min"),
            _spec("SR", "acid_resistance", "class",
                  note="acid resistance SR (ISO method)"),
            _spec("PR", "phosphate_resistance", "class",
                  note="phosphate resistance PR (ISO method)"),
        ],
        "moduli": True, "moduli_gpa": True,
    },
}


def _key(text) -> str:
    """Header text for matching: no spaces or dots, and SCHOTT's acute accent
    ("Young´s") read as an apostrophe. Case counts: "nD" (589.3 nm) sits next
    to "nd" (587.6 nm) in CDGM's and OHARA's sheets."""
    s = str(text or "").replace("’", "'").replace("´", "'")
    # NFKC: HOYA writes codes in full-width letters ("Ｃｐ", "Ｅ(GPa)").
    return re.sub(r"[\s.]", "", unicodedata.normalize("NFKC", s))


def _num(v):
    """A number, or None. Only a cell that IS a number counts: '4-5', '4c',
    '#N/A' and blanks are not values, so nothing is read out of them."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)
    s = str(v).strip()
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def name_key(name) -> str:
    """Glass-name key: the workbooks write 'S-FPM 2' where the .agf has
    'S-FPM2'."""
    return re.sub(r"\s+", "", str(name or "")).upper()


def _header_row(rows, name_header):
    for i, r in enumerate(rows[:10]):
        for j, v in enumerate(r):
            if _key(v) == _key(name_header):
                return i, j
    raise ValueError(f"no {name_header!r} header in the first 10 rows")


def _locate(heads, layout):
    """Column index per spec, by header text, searching the header rows in
    order; first match wins (CDGM repeats 'Glass' at the far right, HOYA's
    'G' is the shear modulus before it is a colour index), except V_d, where
    OHARA prints a rounded value before the exact one."""
    keyed = [[_key(h) for h in head] for head in heads]

    def find(text):
        for keys in keyed:
            hit = [j for j, k in enumerate(keys) if k == _key(text)]
            if hit:
                return hit
        return []
    cols = {}
    for spec in layout["columns"]:
        hit = find(spec["header"])
        if hit:
            cols[spec["prop"]] = (hit[0], spec)
    nd, vd = find(layout["nd"]), find(layout["vd"])
    return cols, (nd[0] if nd else None), (vd[-1] if vd else None)


def _check_schott_units(wb, layout):
    """Refuse the workbook unless its Description sheet still states the
    units the importer assumes."""
    if "Description" not in wb.sheetnames:
        raise ValueError("SCHOTT workbook has no Description sheet")
    desc = {}
    for r in wb["Description"].iter_rows(values_only=True):
        cells = [c for c in r if c is not None]
        for k, c in enumerate(cells[:-1]):
            desc.setdefault(_key(c), str(cells[k + 1]))
    for code, unit in layout["unit_texts"].items():
        text = desc.get(_key(code), "")
        if unit not in text:
            raise ValueError(f"SCHOTT Description for {code!r} no longer "
                             f"states {unit}: {text!r}")


def _check_unit_sheet(wb, layout):
    """Refuse the workbook unless the sheet that prints the units still
    states every one the importer assumes."""
    name = layout["unit_sheet"]
    if name not in wb.sheetnames:
        raise ValueError(f"{layout['maker']} workbook has no {name!r} sheet")
    seen = set()
    for r in wb[name].iter_rows(values_only=True):
        seen.update(_key(c) for c in r if isinstance(c, str))
    for unit in layout["unit_strings"]:
        if not any(_key(unit) in c for c in seen):
            raise ValueError(f"{layout['maker']} {name!r} no longer states "
                             f"{unit}")


def _prop(gid, prop, value, unit, source_id, maker, note):
    return {"glass_id": gid, "property": prop, "value": value, "unit": unit,
            "reference_wavelength_nm": "", "data_type": "manufacturer",
            "source_id": source_id,
            "notes": f"{maker} data workbook: {note}."}


def import_maker_xlsx(path, layout_name: str, source_id: str, names: dict,
                      existing: dict | None = None,
                      catalog_nd_vd: dict | None = None):
    """Read one workbook -> (properties DataFrame, report dict).

    names: name_key(glass name) -> glass_id, for the maker's .agf glasses.
    existing: {(glass_id, property): value} already stored; such properties
      are skipped (a differing value is reported), and rows added here are
      added to it, so a later workbook cannot duplicate them.
    catalog_nd_vd: glass_id -> (n_d, V_d) from the .agf, for the check.
    """
    import openpyxl
    layout = LAYOUTS[layout_name]
    maker = layout["maker"]
    existing = {} if existing is None else existing
    catalog_nd_vd = catalog_nd_vd or {}
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    if layout_name == "SCHOTT":
        _check_schott_units(wb, layout)
    if "unit_sheet" in layout:
        _check_unit_sheet(wb, layout)
    if "sheet_pattern" in layout:
        sheets = [n for n in wb.sheetnames
                  if re.fullmatch(layout["sheet_pattern"], n)]
        if not sheets:
            raise ValueError(f"{maker} workbook has no sheet matching "
                             f"{layout['sheet_pattern']!r}")
    else:
        sheets = layout["sheets"] or wb.sheetnames[:1]
    report = {"rows": 0, "matched": 0, "unmatched": [], "nd_vd": [],
              "conflicts": [], "skipped_existing": 0, "moduli": [],
              "unparsed": 0, "columns": {}}  # columns: sheet -> props it lacks
    out, moduli = [], []
    for sheet in sheets:
        rows = list(wb[sheet].iter_rows(values_only=True))
        hr, name_col = _header_row(rows, layout["name_header"])
        depth = layout.get("header_depth", 1)
        cols, nd_col, vd_col = _locate(rows[hr:hr + depth], layout)
        if "tg" not in cols:
            raise ValueError(f"{maker} sheet {sheet!r}: layout not "
                             f"recognised (no T_g column)")
        report["columns"][sheet] = sorted(
            spec["prop"] for spec in layout["columns"]
            if spec["prop"] not in cols)
        for r in rows[hr + depth:]:
            name = r[name_col] if name_col < len(r) else None
            if not isinstance(name, str) or not name.strip():
                continue
            name = name.strip()
            if name == layout.get("end_marker"):
                break
            report["rows"] += 1
            gid = names.get(name_key(name))
            if gid is None:
                report["unmatched"].append(f"{sheet}: {name}")
                continue
            report["matched"] += 1
            cat = catalog_nd_vd.get(gid)
            if cat:
                for label, col, ref, tol in (("n_d", nd_col, cat[0], ND_TOL),
                                             ("V_d", vd_col, cat[1], VD_TOL)):
                    v = _num(r[col]) if col is not None else None
                    if v is not None and ref is not None and abs(v - ref) > tol:
                        report["nd_vd"].append((gid, label, v, ref))
            vals = {}
            for prop, (j, spec) in cols.items():
                raw = r[j] if j < len(r) else None
                v = _num(raw)
                if v is None:
                    if raw is not None and str(raw).strip() not in ("", "-"):
                        report["unparsed"] += 1
                    continue
                v = round(v * spec["scale"], 6)
                vals[prop] = v
                if (gid, prop) in existing:
                    report["skipped_existing"] += 1
                    old = existing[(gid, prop)]
                    try:
                        same = abs(float(old) - v) < 1e-9
                    except (TypeError, ValueError):
                        same = False
                    if not same:
                        report["conflicts"].append((gid, prop, v, old))
                    continue
                existing[(gid, prop)] = v
                out.append(_prop(gid, prop, v, spec["unit"], source_id, maker,
                                 spec["note"]))
            if layout.get("moduli") and all(
                    k in vals for k in ("youngs_modulus", "shear_modulus",
                                        "poisson_ratio")):
                e, g, nu = (vals["youngs_modulus"], vals["shear_modulus"],
                            vals["poisson_ratio"])
                moduli.append(e)
                if abs(e - 2 * g * (1 + nu)) > MODULI_TOL * e:
                    report["moduli"].append((gid, e, g, nu))
    if moduli:
        if len(report["moduli"]) > MODULI_MAX_FAIL * len(moduli):
            raise ValueError(f"{maker}: E = 2G(1+nu) fails for "
                             f"{len(report['moduli'])} of {len(moduli)} rows - "
                             "wrong column or unit")
        if layout.get("moduli_gpa") and not 40 <= median(moduli) <= 150:
            raise ValueError(f"{maker}: median E {median(moduli)} is not a "
                             "glass modulus in GPa")
    report["moduli_checked"] = len(moduli)
    return pd.DataFrame(out), report
