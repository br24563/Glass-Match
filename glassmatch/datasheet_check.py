"""Cross-check the database's dispersion fits against manufacturers' datasheets.

Every glass curve GlassMatch draws is evaluated from the coefficients in the
maker's Zemax .agf, and the import gate only proves that those coefficients
reproduce the catalog n_d and V_d - three wavelengths. The makers' Excel data
workbooks tabulate the catalog index at 15-20 spectral lines from 365 nm to
2.3 um, written independently of the .agf. This module compares the two at
every printed line the fit's stated range covers, which is the check an
optical engineer would do by hand before trusting a curve.

The workbooks print 5 (OHARA, HOYA) or 6 (CDGM) decimals, so a correct fit
agrees to within rounding plus the fit's own error, a few 1e-6. AGREE_TOL is
deliberately loose enough for 5-decimal tables and tight enough that a wrong
formula, a mis-transcribed coefficient or a different melt shows up.

Pure functions; scripts/validate_against_datasheets.py runs it over the
staged workbooks and writes the committed report.
"""
from __future__ import annotations

import re
import unicodedata

import numpy as np
import pandas as pd

from glassmatch.spectra import n_from_row

# A fit "agrees" with a datasheet when every checked line is within this.
AGREE_TOL = 2e-5

# Spectral lines the workbooks label by letter (um).
LINES_UM = {
    "t": 1.01398, "s": 0.85211, "A'": 0.768195, "r": 0.706519,
    "C": 0.656273, "C'": 0.643847, "He-Ne": 0.632816, "D": 0.589294,
    "d": 0.5875618, "e": 0.546074, "F": 0.4861327, "F'": 0.479992,
    "He-Cd": 0.441569, "g": 0.435835, "h": 0.404656, "i": 0.365015,
}
# IR lines some workbooks abbreviate to whole nanometres ("n2325").
ROUNDED_NM = {2325: 2.32542, 1970: 1.97009, 1530: 1.52958, 1529: 1.52958,
              1129: 1.12864, 1128: 1.12864, 1060: 1.06}


def line_wavelength(label) -> float | None:
    """Wavelength (um) of an index column header ('nC', "nF'", 'n1529.58'),
    or None when the header is not a single-line index (e.g. 'nF-nC')."""
    s = unicodedata.normalize("NFKC", str(label or "")).strip().replace(" ", "")
    m = re.fullmatch(r"n(t|s|A'|r|C|C'|He-Ne|D|d|e|F|F'|He-Cd|g|h|i)", s)
    if m:
        return LINES_UM[m.group(1)]
    m = re.fullmatch(r"n(\d{3,4})(?:\.(\d+))?", s)
    if m:
        nm = float(m.group(0)[1:])
        if m.group(2) is None and int(nm) in ROUNDED_NM:
            return ROUNDED_NM[int(nm)]
        return nm / 1000.0
    return None


def read_index_tables(path, layout_name: str) -> dict:
    """{name_key(glass): {wavelength_um: n}} from one maker workbook.

    Uses the same sheet and header rules as glassmatch.importers.maker_xlsx,
    so it reads exactly the glasses that importer reads."""
    import openpyxl
    from glassmatch.importers.maker_xlsx import (LAYOUTS, _header_row, _num,
                                                  name_key)
    layout = LAYOUTS[layout_name]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "sheet_pattern" in layout:
        sheets = [n for n in wb.sheetnames
                  if re.fullmatch(layout["sheet_pattern"], n)]
    else:
        sheets = layout["sheets"] or wb.sheetnames[:1]
    out: dict = {}
    for sheet in sheets:
        rows = list(wb[sheet].iter_rows(values_only=True))
        hr, name_col = _header_row(rows, layout["name_header"])
        depth = layout.get("header_depth", 1)
        cols = {}
        for head in rows[hr:hr + depth]:
            for j, h in enumerate(head):
                wl = line_wavelength(h)
                if wl is not None and wl not in cols.values():
                    cols[j] = wl
        for r in rows[hr + depth:]:
            name = r[name_col] if name_col < len(r) else None
            if not isinstance(name, str) or not name.strip():
                continue
            if name.strip() == layout.get("end_marker"):
                break
            vals = {wl: _num(r[j]) for j, wl in cols.items() if j < len(r)}
            vals = {wl: v for wl, v in vals.items() if v is not None and v > 1}
            if vals:
                out.setdefault(name_key(name), vals)
    return out


def crosscheck(db, tables: dict, names: dict, source_id: str,
               maker: str) -> pd.DataFrame:
    """One row per glass with both a verified fit and a datasheet table.

    names: name_key -> glass_id for the maker's glasses."""
    from glassmatch.importers.agf import is_dispersable_formula
    rows = []
    for key, table in tables.items():
        gid = names.get(key)
        if gid is None:
            continue
        fit = db.sellmeier_for(gid, sellmeier_only=True)
        if fit is None or not is_dispersable_formula(fit.get("formula")):
            continue
        lo, hi = db.dispersion_range(gid, fit)
        pts = [(wl, n_from_row(fit, wl) - n) for wl, n in sorted(table.items())
               if lo is None or hi is None or lo - 1e-6 <= wl <= hi + 1e-6]
        pts = [(wl, d) for wl, d in pts if d == d]
        if not pts:
            continue
        d = np.array([x for _, x in pts])
        worst = int(np.argmax(np.abs(d)))
        rows.append({"maker": maker, "glass_id": gid, "datasheet": source_id,
                     "lines_printed": len(table), "lines_checked": len(pts),
                     "range_um": f"{pts[0][0]:.4f}-{pts[-1][0]:.4f}",
                     "max_abs_dn": float(np.abs(d).max()),
                     "worst_line_um": pts[worst][0],
                     "rms_dn": float(np.sqrt(np.mean(d * d))),
                     "agrees": bool(np.abs(d).max() <= AGREE_TOL)})
    return pd.DataFrame(rows)
